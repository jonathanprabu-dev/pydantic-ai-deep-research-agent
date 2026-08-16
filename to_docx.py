"""Convert a report markdown file to Word: python to_docx.py all-reports.md

The reports use a small, known subset of markdown — headings, bold claims,
bullets, evidence links and horizontal rules — so this converts that subset
faithfully rather than pulling in a general markdown engine.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET = re.compile(r"^[-*]\s+(.*)$")
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")
RULE = re.compile(r"^-{3,}$")
# Bold, links and inline code, in one pass so they can be interleaved.
INLINE = re.compile(r"\*\*(.+?)\*\*|\[(.*?)\]\((.*?)\)|`(.+?)`")

LINK_BLUE = RGBColor(0x0B, 0x57, 0xD0)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """python-docx has no hyperlink API; the relationship must be added by hand."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = paragraph.add_run(text)
    run.font.color.rgb = LINK_BLUE
    run.font.underline = True
    link.append(run._r)
    paragraph._p.append(link)


def add_rich_text(paragraph, text: str) -> None:
    """Write text, honouring **bold**, [links](url) and `code`."""
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        bold, label, url, code = match.groups()
        if bold is not None:
            paragraph.add_run(bold).bold = True
        elif url is not None:
            add_hyperlink(paragraph, label or url, url)
        elif code is not None:
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def convert(markdown: Path, out: Path) -> Path:
    document = Document()
    document.styles["Normal"].font.size = Pt(10.5)
    first_heading = True

    for line in markdown.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()

        if not line:
            continue

        if RULE.match(line):
            # A rule separates reports in the combined document: start a new page.
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue

        if heading := HEADING.match(line):
            level = len(heading.group(1))
            # Each report opens with "# ...". Keep the first as the title and let
            # the rest be level-1 headings so the Word navigation pane is usable.
            if level == 1 and first_heading:
                document.add_heading(heading.group(2), 0)
                first_heading = False
            else:
                add_rich_text(document.add_heading("", min(level, 4)), heading.group(2))
            continue

        if bullet := BULLET.match(line):
            add_rich_text(document.add_paragraph(style="List Bullet"), bullet.group(1))
            continue

        if numbered := NUMBERED.match(line):
            add_rich_text(document.add_paragraph(style="List Number"), numbered.group(2))
            continue

        add_rich_text(document.add_paragraph(), line)

    document.save(out)
    return out


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    source = Path(sys.argv[1] if len(sys.argv) > 1 else "all-reports.md")
    if not source.is_file():
        raise SystemExit(f"No such file: {source}")
    written = convert(source, source.with_suffix(".docx"))
    print(f"Wrote {written} — {written.stat().st_size:,} bytes")
